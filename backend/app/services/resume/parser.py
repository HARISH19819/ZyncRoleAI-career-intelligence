import os
import re
import pymupdf  # PyMuPDF
import docx
from typing import Dict, Any, List, Optional
from app.agents.job_normalization_agent import normalization_agent
from app.agents.skill_extraction_agent import skill_extraction_agent


class ResumeParserService:
    """Extracts raw text, sections, skills, and metadata from PDF, DOCX, and TXT resumes."""

    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
    MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB

    def extract_text(self, file_path: str, file_ext: str) -> str:
        """Extracts plain text from the uploaded resume file."""
        ext = file_ext.lower()
        if ext == ".pdf":
            text_parts = []
            with pymupdf.open(file_path) as doc:
                for page in doc:
                    text_parts.append(page.get_text())
            return "\n".join(text_parts)
        elif ext == ".docx":
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs)
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    def extract_contact_info(self, text: str) -> Dict[str, Optional[str]]:
        """Extracts email, phone, LinkedIn, GitHub, and portfolio links."""
        email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
        phone_match = re.search(r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text)
        github_match = re.search(r'(?:https?:\/\/)?(?:www\.)?github\.com\/[a-zA-Z0-9_-]+', text, re.IGNORECASE)
        linkedin_match = re.search(r'(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/[a-zA-Z0-9_-]+', text, re.IGNORECASE)
        portfolio_match = re.search(r'(?:https?:\/\/)?(?:www\.)?[a-zA-Z0-9-]+\.(?:dev|me|io|tech|app)', text, re.IGNORECASE)

        return {
            "email": email_match.group(0) if email_match else None,
            "phone": phone_match.group(0) if phone_match else None,
            "github": github_match.group(0) if github_match else None,
            "linkedin": linkedin_match.group(0) if linkedin_match else None,
            "portfolio": portfolio_match.group(0) if portfolio_match else None,
        }

    def detect_sections(self, text: str) -> Dict[str, str]:
        """Segments resume into standard sections."""
        section_headers = [
            "summary", "objective", "professional summary",
            "experience", "work experience", "employment history", "internships",
            "education", "academic background", "academics",
            "projects", "academic projects", "key projects", "personal projects",
            "skills", "technical skills", "core competencies",
            "certifications", "licenses", "awards", "achievements"
        ]

        # Find header positions
        lines = text.split("\n")
        sections: Dict[str, List[str]] = {"general": []}
        current_section = "general"

        for line in lines:
            trimmed = line.strip().lower()
            clean_hdr = re.sub(r'[^a-zA-Z\s]', '', trimmed).strip()
            
            matched_header = None
            if len(clean_hdr) < 35:
                for sh in section_headers:
                    if clean_hdr == sh or clean_hdr.startswith(sh + ":"):
                        matched_header = sh
                        break

            if matched_header:
                # Group related headers
                if any(x in matched_header for x in ["summary", "objective"]):
                    current_section = "summary"
                elif any(x in matched_header for x in ["experience", "employment", "internship"]):
                    current_section = "experience"
                elif any(x in matched_header for x in ["education", "academic"]):
                    current_section = "education"
                elif "project" in matched_header:
                    current_section = "projects"
                elif "skill" in matched_header:
                    current_section = "skills"
                elif any(x in matched_header for x in ["certif", "award", "achievement"]):
                    current_section = "certifications"
                else:
                    current_section = matched_header
                if current_section not in sections:
                    sections[current_section] = []
            else:
                sections[current_section].append(line)

        return {k: "\n".join(v).strip() for k, v in sections.items() if v}

    def extract_quantified_achievements(self, text: str) -> List[str]:
        """Detects bullets with measurable achievements (%, $, numbers, improvement metrics)."""
        lines = text.split("\n")
        achievements = []
        quant_patterns = [
            r'\b\d+%\b',
            r'\bimproved\b.*?\b\d+',
            r'\bincreased\b.*?\b\d+',
            r'\breduced\b.*?\b\d+',
            r'\bscaled\b.*?\b\d+',
            r'\b\d+\s*(?:k|m|million|users|clients|stars)\b',
            r'\$\d+',
        ]
        for line in lines:
            line_str = line.strip()
            if len(line_str) > 20 and len(line_str) < 250:
                for pat in quant_patterns:
                    if re.search(pat, line_str, re.IGNORECASE):
                        achievements.append(line_str.lstrip("•-* "))
                        break
        return achievements[:6]

    def parse(self, file_path: str, file_name: str) -> Dict[str, Any]:
        """Full parsing pipeline producing structured data from a resume file."""
        ext = os.path.splitext(file_name)[1].lower()
        if ext not in self.ALLOWED_EXTENSIONS:
            raise ValueError(f"Extension {ext} is not allowed. Only PDF, DOCX, and TXT are supported.")

        file_size = os.path.getsize(file_path)
        if file_size > self.MAX_FILE_SIZE:
            raise ValueError("File exceeds maximum allowed size of 5 MB.")

        raw_text = self.extract_text(file_path, ext)
        if not raw_text or len(raw_text.strip()) < 50:
            raise ValueError("Resume file appears to be empty or unreadable.")

        contact_info = self.extract_contact_info(raw_text)
        sections = self.detect_sections(raw_text)
        detected_skills = skill_extraction_agent.extract_skills_from_text(raw_text)
        achievements = self.extract_quantified_achievements(raw_text)

        # Detect projects structure
        projects_text = sections.get("projects", "")
        detected_projects = []
        if projects_text:
            p_lines = [l.strip() for l in projects_text.split("\n") if len(l.strip()) > 15]
            # Group into project chunks
            for p in p_lines[:5]:
                detected_projects.append({"title": p[:60], "description": p})

        # Detect experience structure
        exp_text = sections.get("experience", "")
        detected_experience = []
        if exp_text:
            exp_lines = [l.strip() for l in exp_text.split("\n") if len(l.strip()) > 20]
            for e in exp_lines[:4]:
                detected_experience.append({"role": e[:50], "description": e})

        # Detect education
        edu_text = sections.get("education", "")
        detected_education = []
        if edu_text:
            edu_lines = [l.strip() for l in edu_text.split("\n") if len(l.strip()) > 10]
            for ed in edu_lines[:3]:
                detected_education.append({"degree": ed[:80], "institution": ed})

        return {
            "raw_text": raw_text,
            "file_size": file_size,
            "contact_info": contact_info,
            "sections": sections,
            "detected_skills": detected_skills,
            "detected_projects": detected_projects,
            "detected_experience": detected_experience,
            "detected_education": detected_education,
            "achievements": achievements,
        }


resume_parser = ResumeParserService()
